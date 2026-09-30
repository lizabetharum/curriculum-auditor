"""Uploads: Word splits at headings, PDF splits by page with a note, bad files are explained."""
import io

import docx
import pytest

from curriculum_auditor.ingest import PDF_NOTE, load_upload
from curriculum_auditor.segment import SegmentError


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Measuring heights", level=1)
    d.add_paragraph("A one-period lesson.")
    d.add_heading("Task A: The ramp", level=2)
    d.add_paragraph("Use a trig ratio to find the height.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Answer"
    table.cell(1, 0).text, table.cell(1, 1).text = "A3", "0.31 m"
    d.add_heading("Answer key", level=2)
    d.add_paragraph("A3: height = 3.6 × sin(5°) = 0.31 m")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        content = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_word_splits_at_headings_and_keeps_tables():
    result = load_upload("Heights Lesson.docx", make_docx())
    ids = [s.id for s in result.curriculum.sections]
    assert ids == ["measuring-heights", "measuring-heights/task-a-the-ramp", "measuring-heights/answer-key"]
    task = result.curriculum.section("measuring-heights/task-a-the-ramp").text
    assert "A3 | 0.31 m" in task and result.note is None
    assert result.curriculum.name == "Heights Lesson"


def test_pdf_splits_by_page_with_a_note():
    result = load_upload("lesson.pdf", make_pdf(["Warm up with a ramp.", "Answer key: 0.31 m"]))
    assert [s.heading for s in result.curriculum.sections] == ["Page 1", "Page 2"]
    assert "Answer key" in result.curriculum.sections[1].text
    assert result.note == PDF_NOTE


def test_markdown_and_text_still_work():
    assert len(load_upload("l.md", b"# A\ntext\n# B\nmore\n").curriculum.sections) == 2
    assert load_upload("l.txt", b"Just one block of text.").curriculum.sections[0].id == "preamble"


@pytest.mark.parametrize("name,data,message", [
    ("slides.pptx", b"x", "Word"),
    ("empty.md", b"   ", "no text"),
    ("big.md", b"x" * 10_000_001, "10 MB"),
])
def test_bad_uploads_are_explained(name, data, message):
    with pytest.raises(SegmentError, match=message):
        load_upload(name, data)


def test_unreadable_pdf_is_explained():
    with pytest.raises(Exception):
        load_upload("scan.pdf", b"%PDF-1.4 not really")
